"""Inlasning av Swish-rapport (nytt format, en eller flera flikar).

Robust mot formatandringar: rubrikraden hittas dynamiskt via "Bokforingsdatum",
summerings-/tom-/totalrader filtreras bort pa att de saknar bokforingsdatum,
och clearing-/kontonummer lases ur metadatablocket per blad.
"""
from __future__ import annotations

import hashlib
import re
from collections import defaultdict
from datetime import date, datetime, time
from pathlib import Path

import openpyxl

from app.core.models import Swishrapport, Transaktion
from app.core.normalize import to_decimal

_HEADER_MARK = "bokföringsdatum"
_INTERVALL = re.compile(r"(\d{4}-\d{2}-\d{2})\s+to\s+(\d{4}-\d{2}-\d{2})")
_VISNINGSINTERVALL = re.compile(r"(\d{4}-\d{2}-\d{2})\s+-\s+(\d{4}-\d{2}-\d{2})")


class RapportFormatFel(ValueError):
    """Rubrikrad eller kolumner kunde inte kannas igen - begripligt fel."""


def _as_date(v) -> date | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    try:
        return datetime.strptime(s[:10], "%Y-%m-%d").date()
    except ValueError:
        return None


def _as_time(v) -> time | None:
    if v is None or v == "":
        return None
    if isinstance(v, time):
        return v
    if isinstance(v, datetime):
        return v.time()
    s = str(v).strip()
    for fmt in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(s, fmt).time()
        except ValueError:
            continue
    return None


def _text(v) -> str:
    return "" if v is None else str(v).strip()


def _las_metadata(rader: list[list]) -> tuple[str | None, str | None, str | None]:
    """Returnerar (clearingnummer, kontonummer, datumintervall) ur toppblocket."""
    clnr = kontonr = intervall = None
    for rad in rader[:8]:
        for cell in rad:
            t = _text(cell)
            low = t.casefold()
            if low.startswith("clearingnummer:"):
                clnr = t.split(":", 1)[1].strip() or None
            elif low.startswith("kontonummer:"):
                kontonr = t.split(":", 1)[1].strip() or None
            else:
                m = _INTERVALL.search(t)
                if m:
                    intervall = f"{m.group(1)} - {m.group(2)}"
    return clnr, kontonr, intervall


def _hitta_rubrik(rader: list[list]) -> tuple[int, dict[str, int]]:
    for i, rad in enumerate(rader):
        for cell in rad:
            if _text(cell).casefold() == _HEADER_MARK:
                kolmap = {
                    _text(c).casefold(): j
                    for j, c in enumerate(rad) if _text(c)
                }
                return i, kolmap
    raise RapportFormatFel(
        "Hittade ingen rubrikrad med 'Bokföringsdatum'. "
        "Rapportformatet kan ha andrats."
    )


def _col(kolmap: dict[str, int], *namn: str) -> int:
    for n in namn:
        if n in kolmap:
            return kolmap[n]
    raise RapportFormatFel(f"Saknar kolumn: {namn[0]}")


def las_swishrapport(sokvag: str | Path) -> Swishrapport:
    sokvag = Path(sokvag)
    wb = openpyxl.load_workbook(sokvag, data_only=True, read_only=True)
    transaktioner: list[Transaktion] = []
    intervall_totalt: str | None = None

    for sn in wb.sheetnames:
        ws = wb[sn]
        rader = [list(r) for r in ws.iter_rows(values_only=True)]
        if not rader:
            continue
        clnr, kontonr, intervall = _las_metadata(rader)
        intervall_totalt = intervall_totalt or intervall
        hdr_i, kolmap = _hitta_rubrik(rader)

        c_bokf = _col(kolmap, "bokföringsdatum")
        c_trans = _col(kolmap, "transaktionsdatum")
        c_valuta = _col(kolmap, "valutadatum")
        c_mnr = _col(kolmap, "mottagarnummer")
        c_mnamn = _col(kolmap, "mottagarnamn")
        c_medd = _col(kolmap, "meddelande")
        c_order = _col(kolmap, "orderreferens")
        c_tid = _col(kolmap, "tid")
        c_belopp = _col(kolmap, "belopp")

        for rad in rader[hdr_i + 1:]:
            bokf = _as_date(rad[c_bokf] if c_bokf < len(rad) else None)
            if bokf is None:
                continue  # summering/tom/total saknar bokforingsdatum

            def g(idx):
                return rad[idx] if idx < len(rad) else None

            transaktioner.append(Transaktion(
                flik=sn,
                bokf_datum=bokf,
                trans_datum=_as_date(g(c_trans)) or bokf,
                valuta_datum=_as_date(g(c_valuta)),
                mottagarnummer=_text(g(c_mnr)),
                mottagarnamn=_text(g(c_mnamn)),
                meddelande=_text(g(c_medd)),
                orderreferens=_text(g(c_order)),
                tid=_as_time(g(c_tid)),
                belopp=to_decimal(g(c_belopp)),
                clearingnummer=clnr,
                kontonummer=kontonr,
            ))
    wb.close()

    _satt_tx_id(transaktioner)
    period = _harled_period(intervall_totalt, transaktioner)
    return Swishrapport(
        filnamn=sokvag.name,
        period=period,
        datumintervall=intervall_totalt or "",
        transaktioner=transaktioner,
    )


def _satt_tx_id(txs: list[Transaktion]) -> None:
    """Stabil identitet per transaktion (overlever omlasning av samma fil).

    Hash av innehallet + ett lopnummer for exakta dubletter (samma tid, belopp
    och meddelande), i filordning.
    """
    seen: dict[str, int] = defaultdict(int)
    for t in txs:
        bas = "|".join([
            t.flik, str(t.bokf_datum), str(t.trans_datum), str(t.tid),
            t.mottagarnummer, str(t.belopp), t.meddelande,
        ])
        n = seen[bas]
        seen[bas] += 1
        t.tx_id = hashlib.sha1(f"{bas}#{n}".encode()).hexdigest()[:12]


def _harled_period(intervall: str | None, txs: list[Transaktion]) -> str:
    """Period 'YYYY-MM' ur intervallets startdatum, annars vanligaste transmanad."""
    if intervall:
        m = _INTERVALL.search(intervall) or _VISNINGSINTERVALL.search(intervall)
        if m:
            return m.group(1)[:7]
    if txs:
        from collections import Counter
        c = Counter(t.trans_datum.strftime("%Y-%m") for t in txs)
        return c.most_common(1)[0][0]
    return ""
