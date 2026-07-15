"""Inlasning av andamalskalendern (originalformat: ett blad per forsamling)."""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import openpyxl

from app.config import Kollekttyp
from app.core.models import Kalenderrad
from app.core.normalize import forsamling_fran_kortkod


def _as_date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if v in (None, ""):
        return None
    try:
        return datetime.strptime(str(v)[:10], "%Y-%m-%d").date()
    except ValueError:
        return None


def _text(v) -> str:
    return "" if v is None else str(v).strip()


def _typ(v) -> Kollekttyp | None:
    t = _text(v).upper()
    return Kollekttyp(t) if t in ("F", "R", "S") else None


def las_kalender(sokvag: str | Path) -> list[Kalenderrad]:
    wb = openpyxl.load_workbook(sokvag, data_only=True, read_only=True)
    rader: list[Kalenderrad] = []

    for sn in wb.sheetnames:
        fors = forsamling_fran_kortkod(sn)
        if fors is None:
            continue  # ej ett forsamlingsblad
        for i, row in enumerate(wb[sn].iter_rows(values_only=True)):
            if i == 0:
                continue  # rubrikrad
            vals = list(row)
            # Kolumnordning: Datum, Tema, Veckodag, Kyrklig helgdag, Typ, Andamal
            datum = _as_date(vals[0] if len(vals) > 0 else None)
            if datum is None:
                continue
            rader.append(Kalenderrad(
                forsamling=fors.kanoniskt,
                kortkod=fors.kortkod,
                datum=datum,
                veckodag=_text(vals[2]) if len(vals) > 2 else "",
                helgdag=_text(vals[3]) if len(vals) > 3 else "",
                typ=_typ(vals[4]) if len(vals) > 4 else None,
                andamal=_text(vals[5]) if len(vals) > 5 else "",
            ))
    wb.close()
    rader.sort(key=lambda r: (r.forsamling, r.datum))
    return rader
